#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
xlsx_to_json.py — data/species.json را کامل از روی اکسل کل دیتابیس
(خروجی build_database_template.py) بازمی‌سازد. از این مرحله به بعد، اکسل
منبع اصلی داده است، نه CSV رسمی MDD — شیت «گونه‌ها» هر فیلدی که دارد را
عیناً به data/species.json می‌برد (ردیف پاک‌شده = گونه حذف‌شده، ردیف
تازه = گونه‌ی جدید)، و شیت «پراکنش_استانی» فهرست پراکنش هر گونه را می‌سازد.

اگر بعداً خواستید داده را از نسخه‌ی تازه‌ی MDD به‌روزرسانی کنید،
scripts/mdd_iran_to_json.py را جدا اجرا کنید و بعد یک اکسل تازه با
scripts/build_database_template.py بسازید — این اسکریپت خودش با MDD تماس
نمی‌گیرد.

استفاده:
    python3 xlsx_to_json.py مسیر/به/دیتابیس_پستانداران_ایران.xlsx [data/species.json]
"""

import json
import sys
from pathlib import Path

from openpyxl import load_workbook

from species_fields import FIELDS, cell_to_field, VALID_IUCN

PROVINCES = {
    "آذربایجان شرقی", "آذربایجان غربی", "اردبیل", "اصفهان", "البرز", "ایلام",
    "بوشهر", "تهران", "چهارمحال و بختیاری", "خراسان جنوبی", "خراسان رضوی",
    "خراسان شمالی", "خوزستان", "زنجان", "سمنان", "سیستان و بلوچستان", "فارس",
    "قزوین", "قم", "کردستان", "کرمان", "کرمانشاه", "کهگیلویه و بویراحمد",
    "گلستان", "گیلان", "لرستان", "مازندران", "مرکزی", "هرمزگان", "همدان", "یزد",
}

PRESENCE_OPTIONS = {
    "حضور دارد", "احتمالی/نیازمند تأیید", "گزارش تاریخی", "منقرض‌شده محلی",
}

EXAMPLE_MARKER = "نمونه"
ID_COL = 1  # ستون «شناسه» در شیت «گونه‌ها» — همیشه اول باید بماند


def die(msg: str) -> None:
    print(f"خطا: {msg}", file=sys.stderr)
    sys.exit(1)


def cell(ws, row, col):
    v = ws.cell(row=row, column=col).value
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def row_is_empty(ws, row, n_cols):
    return all(cell(ws, row, c) is None for c in range(1, n_cols + 1))


def load_species_sheet(ws):
    max_row = ws.max_row  # یک‌بار خوانده می‌شود — دلیل را در load_distribution ببینید
    n_cols = len(FIELDS)
    species = []
    seen_ids = {}
    errors = []
    next_auto_num = 1

    raw_rows = []
    for row in range(2, max_row + 1):
        if row_is_empty(ws, row, n_cols):
            continue
        raw_rows.append(row)

    # اول شناسه‌های صریح را جمع می‌کنیم تا شماره‌ی خودکارِ بعدی با آن‌ها برخورد نکند
    for row in raw_rows:
        sid = cell(ws, row, ID_COL)
        if sid and sid.startswith("IR") and sid[2:].isdigit():
            next_auto_num = max(next_auto_num, int(sid[2:]) + 1)

    for row in raw_rows:
        sp = {}
        for col_i, spec in enumerate(FIELDS, start=1):
            raw = cell(ws, row, col_i)
            cell_to_field(sp, spec, raw)

        if not sp.get("scientificName"):
            errors.append(f"سطر {row} (شیت «گونه‌ها»): نام علمی خالی است")
            continue

        if not sp.get("id"):
            sp["id"] = f"IR{next_auto_num:04d}"
            next_auto_num += 1

        if sp["id"] in seen_ids:
            errors.append(
                f"سطر {row} (شیت «گونه‌ها»): شناسه «{sp['id']}» تکراری است "
                f"(قبلاً در سطر {seen_ids[sp['id']]} دیده شد)"
            )
            continue
        seen_ids[sp["id"]] = row

        # ستون‌های محاسبه‌شده
        sp["mddUrl"] = f"https://www.mammaldiversity.org/taxon/{sp['mddId']}" if sp.get("mddId") else None

        # order/family هم داخل taxonomy تکرار می‌شود (همان شکلی که mdd_iran_to_json.py می‌سازد)
        sp.setdefault("taxonomy", {})
        sp["taxonomy"]["order"] = sp.get("order")
        sp["taxonomy"]["family"] = sp.get("family")

        sp["worldDistribution"] = (sp.get("distribution") or {}).get("countries", [])
        sp["provinces"] = []  # با شیت پراکنش_استانی پر می‌شود

        species.append(sp)

    return species, errors


def load_distribution(ws, valid_ids):
    max_row = ws.max_row  # نکته: ws.cell() با probe کردن ردیف‌های فراتر از انتهای واقعی، خودش
    # سلول خالی می‌سازد و max_row را جلو می‌برد؛ برای همین این مقدار فقط یک‌بار قبل از حلقه
    # خوانده می‌شود، نه در هر تکرار — وگرنه حلقه هرگز تمام نمی‌شود.
    by_species = {}
    errors = []
    skipped_examples = 0

    for row in range(2, max_row + 1):
        sid = cell(ws, row, 1)
        province = cell(ws, row, 3)
        presence = cell(ws, row, 4)
        source = cell(ws, row, 5)
        year = cell(ws, row, 6)
        note = cell(ws, row, 7)

        if note and str(note).startswith(EXAMPLE_MARKER):
            skipped_examples += 1
            continue

        if not any([sid, province, presence, source, year, note]):
            continue

        if sid not in valid_ids:
            errors.append(f"سطر {row} (شیت «پراکنش_استانی»): شناسه‌گونه «{sid}» در شیت «گونه‌ها» نیست")
        if province not in PROVINCES:
            errors.append(f"سطر {row} (شیت «پراکنش_استانی»): استان «{province}» نامعتبر است")
        if presence not in PRESENCE_OPTIONS:
            errors.append(f"سطر {row} (شیت «پراکنش_استانی»): وضعیت حضور «{presence}» نامعتبر است")
        if not source:
            errors.append(f"سطر {row} (شیت «پراکنش_استانی»): ستون «منبع» خالی است — بدون منبع رکورد منتشر نمی‌شود")

        if sid in valid_ids and province in PROVINCES and presence in PRESENCE_OPTIONS and source:
            by_species.setdefault(sid, []).append({
                "name": province, "presence": presence,
                "source": source, "year": year, "note": note,
            })

    return by_species, errors, skipped_examples


def main():
    if len(sys.argv) < 2:
        die("استفاده: python3 xlsx_to_json.py مسیر_اکسل.xlsx [مسیر_species.json]")

    xlsx_path = Path(sys.argv[1])
    json_path = Path(sys.argv[2]) if len(sys.argv) >= 3 else (
        Path(__file__).resolve().parent.parent / "data" / "species.json"
    )

    if not xlsx_path.exists():
        die(f"فایل پیدا نشد: {xlsx_path}")

    wb = load_workbook(xlsx_path, data_only=True)
    for sheet in ("گونه‌ها", "پراکنش_استانی"):
        if sheet not in wb.sheetnames:
            die(f"شیت «{sheet}» در اکسل پیدا نشد. از build_database_template.py استفاده کرده‌اید؟")

    species, sp_errors = load_species_sheet(wb["گونه‌ها"])
    if sp_errors:
        die(f"{len(sp_errors)} خطا در شیت «گونه‌ها»:\n  " + "\n  ".join(sp_errors[:20])
            + ("\n  ..." if len(sp_errors) > 20 else ""))

    valid_ids = {sp["id"] for sp in species}
    dist_by_species, dist_errors, skipped = load_distribution(wb["پراکنش_استانی"], valid_ids)
    if dist_errors:
        die(f"{len(dist_errors)} خطا در شیت «پراکنش_استانی»:\n  " + "\n  ".join(dist_errors[:20])
            + ("\n  ..." if len(dist_errors) > 20 else ""))

    covered = 0
    total_records = 0
    for sp in species:
        provinces = dist_by_species.get(sp["id"], [])
        sp["provinces"] = provinces
        if provinces:
            covered += 1
            total_records += len(provinces)

    meta = {
        "version": "xlsx-full-1",
        "updated": __import__("datetime").date.today().isoformat(),
        "note": (
            f"{len(species)} گونه — منبع اصلی این نسخه خودِ اکسل {xlsx_path.name} است (نه دیگر "
            "مستقیماً CSV رسمی MDD؛ برای به‌روزرسانی از MDD، mdd_iran_to_json.py را جدا اجرا کنید "
            "و یک اکسل تازه بسازید). "
            f"{covered} گونه حداقل یک رکورد پراکنش استانی دارند ({total_records} رکورد)."
        ),
        "source": "اکسل دیتابیس پستانداران ایران (ورودی انسانی، بر پایه‌ی داده‌ی MDD)",
    }

    if json_path.exists():
        try:
            old = json.loads(json_path.read_text(encoding="utf-8"))
            meta = {**old.get("meta", {}), **meta}
        except Exception:
            pass

    json_path.parent.mkdir(parents=True, exist_ok=True)
    json_path.write_text(
        json.dumps({"meta": meta, "species": species}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    print(f"نوشته شد: {json_path}")
    print(f"  {len(species)} گونه، {covered} گونه با پراکنش استانی ({total_records} رکورد)")
    if skipped:
        print(f"  {skipped} ردیف نمونه نادیده گرفته شد")


if __name__ == "__main__":
    main()
