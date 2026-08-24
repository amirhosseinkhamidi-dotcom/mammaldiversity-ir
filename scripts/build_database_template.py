#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_database_template.py — کل data/species.json را به یک اکسل چهارشیتی
صادر می‌کند: یک ردیف اکسل = یک گونه، با تقریباً همه‌ی فیلدهای دیتاست (نه فقط
پراکنش استانی). این اکسل خودِ منبع اصلی داده می‌شود — هر فیلدی را که لازم شد
اصلاح/اضافه/حذف کنید، مستقیم در همین‌جا انجام می‌شود.

هر بار که چیزی در اکسل عوض شد:
    python3 scripts/xlsx_to_json.py data/دیتابیس_پستانداران_ایران.xlsx

این دستور data/species.json را از روی محتوای فعلیِ اکسل بازمی‌سازد (کل
فهرست گونه‌ها را جایگزین می‌کند — اضافه/حذف ردیف هم همین‌جا اثر می‌گذارد).

اگر خواستید دوباره از اکسلِ فعلی + هر تغییری که در data/species.json افتاده
(مثلاً بعد از بازتولید از نسخه‌ی جدید MDD) یک اکسل تازه بسازید، همین اسکریپت
را دوباره اجرا کنید — فایل قبلی بازنویسی می‌شود.

استفاده:
    python3 build_database_template.py [مسیر_خروجی.xlsx]
"""

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

from species_fields import FIELDS, field_to_cell

ROOT = Path(__file__).resolve().parent.parent
FONT_NAME = "Tahoma"  # بهترین پشتیبانی از حروف فارسی در ویندوز/اکسل

PROVINCES = [
    "آذربایجان شرقی", "آذربایجان غربی", "اردبیل", "اصفهان", "البرز", "ایلام",
    "بوشهر", "تهران", "چهارمحال و بختیاری", "خراسان جنوبی", "خراسان رضوی",
    "خراسان شمالی", "خوزستان", "زنجان", "سمنان", "سیستان و بلوچستان", "فارس",
    "قزوین", "قم", "کردستان", "کرمان", "کرمانشاه", "کهگیلویه و بویراحمد",
    "گلستان", "گیلان", "لرستان", "مازندران", "مرکزی", "هرمزگان", "همدان", "یزد",
]

PRESENCE_OPTIONS = [
    "حضور دارد",
    "احتمالی/نیازمند تأیید",
    "گزارش تاریخی",
    "منقرض‌شده محلی",
]

HEADER_FILL = PatternFill("solid", fgColor="0F766E")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)
INPUT_FILL = PatternFill("solid", fgColor="FFF6D8")
COMPUTED_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top", horizontal="right", readingOrder=2)
CENTER = Alignment(horizontal="center", vertical="center", readingOrder=2)


def style_sheet_rtl(ws):
    ws.sheet_view.rightToLeft = True


def header_row(ws, headers, row=1, widths=None):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = CENTER
        c.border = BORDER
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"A{row + 1}"


def build_instructions_sheet(wb, n_species):
    ws = wb.active
    ws.title = "راهنما"
    style_sheet_rtl(ws)
    ws.column_dimensions["A"].width = 100

    lines = [
        ("پایگاه پستانداران ایران — اکسل کامل دیتابیس", True, 15),
        ("", False, 11),
        (f"این فایل کل دیتاست فعلی سایت را دارد ({n_species} گونه، تقریباً همه‌ی فیلدهایی که "
         "در صفحه‌ی هر گونه نشان داده می‌شود). برخلاف نسخه‌های قبلی این پروژه، از این به بعد "
         "خودِ همین اکسل منبع اصلی داده است — هر فیلدی که لازم شد اصلاح، اضافه یا حذف کنید، "
         "مستقیم این‌جا انجام می‌شود، بعد با اسکریپت پایین به data/species.json برمی‌گردد.", False, 11),
        ("", False, 11),
        ("شیت «گونه‌ها»", True, 13),
        ("یک ردیف = یک گونه. تقریباً همه‌ی ستون‌ها زرد و قابل‌ویرایش‌اند. دو استثنا:", False, 11),
        ("  • «پیوند MDD» (خاکستری): همیشه خودکار از روی «شناسه MDD» ساخته می‌شود؛ هرچه این‌جا "
         "بنویسید در ادغام بعدی نادیده گرفته می‌شود.", False, 11),
        ("  • «شناسه»: کلید یکتای هر گونه (IRxxxx) است. برای ردیف‌های موجود دست نزنید — شیت "
         "«پراکنش_استانی» با همین شناسه به گونه وصل می‌شود. برای گونه‌ی تازه‌ای که خودتان اضافه "
         "می‌کنید، این ستون را خالی بگذارید؛ خودکار یک شناسه‌ی جدید می‌گیرد.", False, 11),
        ("برای حذف یک گونه از دیتاست، کل ردیفش را پاک کنید.", False, 11),
        ("ستون‌هایی که چند مقدار دارند (نام‌های دیگر، نام‌های مترادف، کشورها، قاره‌ها، ...) با "
         "علامت « | » از هم جدا می‌شوند — مثلاً: Iran | Iraq | Turkey", False, 11),
        ("ستون‌های بله/خیر (بومی، منقرض‌شده، اهلی، نیازمند بازبینی، تفاوت با CMW/MSW3) را دقیقاً "
         "با «بله» یا «خیر» پر کنید.", False, 11),
        ("ستون «وضعیت IUCN» باید یکی از این کدها باشد: CR, EN, VU, NT, LC, DD, NE, EX, EW", False, 11),
        ("", False, 11),
        ("شیت «پراکنش_استانی»", True, 13),
        ("هر ردیف یعنی «این گونه در این استان با این وضعیت دیده/گزارش شده». برای هر گونه می‌توانید "
         "چند ردیف اضافه کنید. ستون «منبع» الزامی است — بدون منبع، رکورد منتشر نمی‌شود.", False, 11),
        ("  • شناسه‌گونه: از فهرست کشویی (همان ستون «شناسه» در شیت گونه‌ها).", False, 11),
        ("  • نام علمی: خودکار پر می‌شود؛ ویرایش نکنید.", False, 11),
        ("  • استان: از فهرست کشویی.", False, 11),
        ("  • وضعیت حضور: حضور دارد / احتمالی-نیازمند تأیید / گزارش تاریخی / منقرض‌شده محلی.", False, 11),
        ("  • منبع: الزامی — نام مقاله/گزارش+سال، مشاهدهٔ میدانی+تیم+سال، تله‌عکس+محل+سال، یا هر "
         "منبع قابل‌پیگیری دیگر.", False, 11),
        ("  • سال ثبت و یادداشت: اختیاری.", False, 11),
        ("", False, 11),
        ("شیت «استان‌ها»", True, 13),
        ("فقط فهرست ۳۱ استان برای فهرست‌های کشویی؛ نیازی به ویرایش ندارد.", False, 11),
        ("", False, 11),
        ("بعد از هر تغییر", True, 13),
        ("فایل را ذخیره کنید و اجرا کنید:", False, 11),
        ("python3 scripts/xlsx_to_json.py data/دیتابیس_پستانداران_ایران.xlsx", False, 11),
        ("این دستور data/species.json را کامل از روی محتوای همین اکسل بازمی‌سازد — اگر ردیفی را "
         "پاک کرده باشید، آن گونه از سایت هم حذف می‌شود؛ اگر ردیف تازه اضافه کرده باشید، به سایت "
         "اضافه می‌شود. اگر جایی خطا داشته باشد (شناسه تکراری، IUCN نامعتبر، منبع خالی در شیت "
         "پراکنش)، اسکریپت متوقف می‌شود و دقیقاً می‌گوید کجا را اصلاح کنید.", False, 11),
    ]

    r = 1
    for text, bold, size in lines:
        c = ws.cell(row=r, column=1, value=text)
        c.font = Font(name=FONT_NAME, bold=bold, size=size, color="0F766E" if bold and size >= 13 else "17231F")
        c.alignment = WRAP
        if text:
            ws.row_dimensions[r].height = 20 if not bold else (28 if size >= 15 else 22)
        r += 1


def build_species_sheet(wb, species):
    ws = wb.create_sheet("گونه‌ها")
    style_sheet_rtl(ws)
    headers = [f.header for f in FIELDS]
    widths = [f.width for f in FIELDS]
    header_row(ws, headers, widths=widths)

    for spec_i, spec in enumerate(FIELDS):
        if not spec.editable:
            ws.cell(row=1, column=spec_i + 1).comment = Comment(
                "محاسبه‌شده — ویرایش نکنید (خودکار از «شناسه MDD» ساخته می‌شود).", FONT_NAME
            )

    for row_i, sp in enumerate(species, start=2):
        for col_i, spec in enumerate(FIELDS, start=1):
            val = field_to_cell(sp, spec)
            c = ws.cell(row=row_i, column=col_i, value=val)
            c.font = Font(name=FONT_NAME, size=10, italic=spec.italic)
            c.border = BORDER
            c.alignment = WRAP
            c.fill = INPUT_FILL if spec.editable else COMPUTED_FILL

    # فهرست کشویی برای ستون «وضعیت IUCN»
    iucn_col = next(i for i, f in enumerate(FIELDS, start=1) if f.path == "iucn")
    dv_iucn = DataValidation(type="list", formula1='"CR,EN,VU,NT,LC,DD,NE,EX,EW"', allow_blank=True)
    ws.add_data_validation(dv_iucn)
    dv_iucn.add(f"{get_column_letter(iucn_col)}2:{get_column_letter(iucn_col)}{len(species) + 400}")

    # فهرست کشویی برای ستون‌های بله/خیر
    bool_cols = [i for i, f in enumerate(FIELDS, start=1) if f.kind == "bool"]
    dv_bool = DataValidation(type="list", formula1='"بله,خیر"', allow_blank=True)
    ws.add_data_validation(dv_bool)
    for col in bool_cols:
        dv_bool.add(f"{get_column_letter(col)}2:{get_column_letter(col)}{len(species) + 400}")

    return len(species)


def build_provinces_sheet(wb):
    ws = wb.create_sheet("استان‌ها")
    style_sheet_rtl(ws)
    header_row(ws, ["نام استان"], widths=[26])
    for i, p in enumerate(PROVINCES, start=2):
        c = ws.cell(row=i, column=1, value=p)
        c.font = Font(name=FONT_NAME, size=10)
        c.border = BORDER
        c.alignment = WRAP
    return len(PROVINCES)


def build_distribution_sheet(wb, species, n_species, n_provinces):
    ws = wb.create_sheet("پراکنش_استانی")
    style_sheet_rtl(ws)
    headers = ["شناسه‌گونه", "نام علمی", "استان", "وضعیت حضور", "منبع", "سال ثبت", "یادداشت"]
    header_row(ws, headers, widths=[13, 26, 20, 22, 34, 10, 30])

    # هر رکورد پراکنشی که از قبل در species.json بود را عیناً برمی‌گردانیم (round-trip بی‌اتلاف).
    existing_rows = []
    for sp in species:
        for pr in sp.get("provinces", []) or []:
            existing_rows.append((
                sp["id"], pr.get("name"), pr.get("presence"),
                pr.get("source"), pr.get("year"), pr.get("note"),
            ))

    if not existing_rows:
        existing_rows = [
            ("IR0046", "مازندران", "حضور دارد", "پایش دوربین‌تله‌ای پروژهٔ حفاظت پلنگ، ۱۴۰۲", "1402",
             "نمونه — این ردیف را با داده‌ی واقعی جایگزین یا پاک کنید"),
            ("IR0046", "گلستان", "احتمالی/نیازمند تأیید", "گزارش شفاهی محیط‌بانان، تأییدنشده", "1401",
             "نمونه — این ردیف را هم پاک یا جایگزین کنید"),
        ]

    LAST_ROW = max(800, len(existing_rows) + 50)

    for i, row_data in enumerate(existing_rows, start=2):
        sid, province, presence, source, year, note = row_data
        vals = [sid, None, province, presence, source, year, note]
        for col, val in enumerate(vals, start=1):
            c = ws.cell(row=i, column=col, value=val)
            c.font = Font(name=FONT_NAME, size=10, italic=(col == 2))
            c.border = BORDER
            c.alignment = WRAP
            if col != 2:
                c.fill = INPUT_FILL

    for i in range(2, LAST_ROW + 1):
        formula = (
            f"=IFERROR(INDEX('گونه‌ها'!$G$2:$G${n_species + 1},"
            f"MATCH(A{i},'گونه‌ها'!$A$2:$A${n_species + 1},0)),\"\")"
        )
        c = ws.cell(row=i, column=2, value=formula)
        c.font = Font(name=FONT_NAME, size=10, italic=True, color="6E7D78")
        c.border = BORDER
        c.alignment = WRAP
        for col in (1, 3, 4, 5, 6, 7):
            ws.cell(row=i, column=col).border = BORDER

    dv_id = DataValidation(type="list", formula1=f"='گونه‌ها'!$A$2:$A${n_species + 1}",
                            allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv_id)
    dv_id.add(f"A2:A{LAST_ROW}")

    dv_prov = DataValidation(type="list", formula1=f"='استان‌ها'!$A$2:$A${n_provinces + 1}",
                              allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv_prov)
    dv_prov.add(f"C2:C{LAST_ROW}")

    dv_presence = DataValidation(type="list", formula1='"' + ",".join(PRESENCE_OPTIONS) + '"',
                                  allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv_presence)
    dv_presence.add(f"D2:D{LAST_ROW}")

    ws["B1"].comment = Comment(
        "خودکار پر می‌شود — ویرایش نکنید. برای راهنمای کامل به شیت «راهنما» نگاه کنید.", FONT_NAME
    )


def main():
    out_path = Path(sys.argv[1]) if len(sys.argv) > 1 else (
        ROOT / "data" / "دیتابیس_پستانداران_ایران.xlsx"
    )
    species = json.loads((ROOT / "data" / "species.json").read_text(encoding="utf-8"))["species"]

    # ستون هفتم شیت «گونه‌ها» باید نام علمی باشد تا فرمول شیت پراکنش درست کار کند
    assert FIELDS[6].path == "scientificName", "ترتیب FIELDS عوض شده؛ فرمول ستون B را در build_distribution_sheet هم به‌روزرسانی کنید"

    wb = Workbook()
    build_instructions_sheet(wb, len(species))
    n_species = build_species_sheet(wb, species)
    n_provinces = build_provinces_sheet(wb)
    build_distribution_sheet(wb, species, n_species, n_provinces)

    wb.active = 1
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"نوشته شد: {out_path}")
    print(f"  {n_species} گونه، {sum(len(sp.get('provinces', []) or []) for sp in species)} رکورد پراکنش استانی موجود")


if __name__ == "__main__":
    main()
